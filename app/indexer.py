"""Local document ingestion and SQLite FTS5 index. Never copies documents to Git."""
from __future__ import annotations
import hashlib, html, json, re, sqlite3
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED={'.md','.txt','.html','.htm','.docx','.xlsx','.xls','.xlsm','.csv','.json'}
class HTMLText(HTMLParser):
    def __init__(self): super().__init__(); self.text=[]; self.ignore=0
    def handle_starttag(self,tag,attrs):
        if tag in ('script','style','noscript'):self.ignore+=1
        if tag in ('p','div','li','h1','h2','h3','tr','br'):self.text.append('\n')
    def handle_endtag(self,tag):
        if tag in ('script','style','noscript'):self.ignore=max(0,self.ignore-1)
    def handle_data(self,data):
        if not self.ignore:self.text.append(data)

def extract(p:Path)->str:
    ext=p.suffix.lower()
    if ext in ('.md','.txt','.csv'): return p.read_text(encoding='utf-8-sig',errors='replace')
    if ext in ('.html','.htm'):
        h=HTMLText();h.feed(p.read_text(encoding='utf-8-sig',errors='replace'));return html.unescape(' '.join(h.text))
    if ext=='.json':return json.dumps(json.loads(p.read_text(encoding='utf-8-sig')),ensure_ascii=False)
    if ext=='.docx':
        from docx import Document
        d=Document(str(p))
        parts=[x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows: parts.append(' | '.join(c.text for c in row.cells))
        return '\n'.join(parts)
    if ext in ('.xlsx','.xls','.xlsm'):
        from python_calamine import CalamineWorkbook
        wb=CalamineWorkbook.from_path(str(p))
        rows=[]
        for name in wb.sheets:
            sheet=wb.get_sheet_by_name(name)
            rows.append('SHEET: '+name)
            for row in sheet.to_python():rows.append(' | '.join(str(x) if x is not None else '' for x in row))
        return '\n'.join(rows)
    raise ValueError('Unsupported extension')

def split_text(text:str,chunk_chars:int=1200,overlap:int=160):
    text=re.sub(r'\s+',' ',text).strip()
    if not text:return []
    parts=[];start=0
    while start<len(text):
        end=min(len(text),start+chunk_chars)
        if end<len(text):
            cut=text.rfind('. ',start+chunk_chars//2,end)
            if cut>start: end=cut+1
        parts.append(text[start:end])
        if end==len(text):break
        start=max(start+1,end-overlap)
    return parts

def open_index(index_path:Path):
    index_path.parent.mkdir(parents=True,exist_ok=True)
    c=sqlite3.connect(str(index_path),timeout=20)
    c.execute('PRAGMA journal_mode=WAL')
    c.execute('CREATE TABLE IF NOT EXISTS documents(path TEXT PRIMARY KEY,sha256 TEXT NOT NULL,mtime REAL NOT NULL,chunk_count INTEGER NOT NULL)')
    c.execute('CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5(path UNINDEXED,chunk_id UNINDEXED,body,tokenize="unicode61 remove_diacritics 2")')
    return c

def build_index(folder:Path,index_path:Path,max_file_mb:int=10,max_files:int=1000):
    folder=folder.resolve()
    if not folder.is_dir():raise ValueError('Knowledge directory missing')
    c=open_index(index_path);seen=set();counts={'indexed':0,'unchanged':0,'removed':0,'skipped':0,'errors':[]}
    try:
        candidates=(p for p in sorted(folder.rglob('*')) if p.is_file() and not p.is_symlink() and p.suffix.lower() in SUPPORTED)
        for n,p in enumerate(candidates):
            if n>=max_files:raise ValueError('Too many documents')
            path=str(p.resolve());seen.add(path)
            if p.stat().st_size>max_file_mb*1024*1024:
                counts['skipped']+=1;counts['errors'].append({'file':p.name,'reason':'too large'});continue
            try:
                digest=hashlib.sha256(p.read_bytes()).hexdigest()
                old=c.execute('SELECT sha256 FROM documents WHERE path=?',(path,)).fetchone()
                if old and old[0]==digest:counts['unchanged']+=1;continue
                text=extract(p);chunks=split_text(text)
                with c:
                    c.execute('DELETE FROM chunks WHERE path=?',(path,))
                    for i,part in enumerate(chunks):c.execute('INSERT INTO chunks(path,chunk_id,body) VALUES (?,?,?)',(path,str(i),part))
                    c.execute('INSERT OR REPLACE INTO documents VALUES (?,?,?,?)',(path,digest,p.stat().st_mtime,len(chunks)))
                counts['indexed']+=1
            except Exception as e:
                counts['errors'].append({'file':p.name,'reason':type(e).__name__})
        for (path,) in c.execute('SELECT path FROM documents').fetchall():
            if path not in seen:
                with c:
                    c.execute('DELETE FROM chunks WHERE path=?',(path,))
                    c.execute('DELETE FROM documents WHERE path=?',(path,))
                counts['removed']+=1
        counts['documents']=c.execute('SELECT COUNT(*) FROM documents').fetchone()[0]
        counts['chunks']=c.execute('SELECT COUNT(*) FROM chunks').fetchone()[0]
        return counts
    finally:c.close()

def search_index(index_path:Path,question:str,top_k:int=4):
    if not index_path.is_file():return []
    terms=re.findall(r'[^\W_]{3,}',question,flags=re.UNICODE)
    if not terms:return []
    query=' OR '.join('"'+s.replace('"','')+'"' for s in list(dict.fromkeys(terms))[:12])
    c=open_index(index_path)
    try:
        rows=c.execute('SELECT path,chunk_id,body FROM chunks WHERE chunks MATCH ? ORDER BY bm25(chunks) LIMIT ?',(query,top_k)).fetchall()
        return [{'id':Path(p).name+'#'+i,'source':p,'text':body} for p,i,body in rows]
    finally:c.close()
